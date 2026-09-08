# -*- coding: utf-8 -*-
"""
📊 VegaWin Dosya, Fiyat ve Toplu Stok Senkronizasyon Servisi
"""
import os
import csv
import datetime
import openpyxl
from backend.services.db_service import db_session, init_db, clean_barcode_text, record_product_history

def normalize_text(text: str) -> str:
    return str(text).strip() if text is not None else ""

def parse_price(val) -> float:
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    s = str(val).strip().replace('TL', '').replace('tl', '').replace('₺', '').strip()
    if ',' in s and '.' in s:
        last_comma = s.rfind(',')
        last_dot = s.rfind('.')
        if last_comma > last_dot:
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        s = s.replace(',', '.')
    try:
        return round(float(s), 2)
    except Exception:
        return 0.0

def parse_vegawin_file(file_path: str) -> list:
    ext = os.path.splitext(file_path)[1].lower()
    raw_rows = []

    if ext in ['.xlsx', '.xls']:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        sheet = wb.active
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None and str(cell).strip() != '' for cell in row):
                raw_rows.append(list(row))
    elif ext == '.csv':
        encodings = ['utf-8-sig', 'utf-8', 'cp1254', 'windows-1254', 'iso-8859-9', 'latin-1']
        content = None
        for enc in encodings:
            try:
                with open(file_path, 'r', encoding=enc) as f:
                    content = f.read()
                    break
            except Exception:
                continue
        if not content:
            raise ValueError("CSV dosyası okunamadı veya boş.")

        sample_lines = [l for l in content.splitlines() if l.strip()][:10]
        delimiter = ';' if sum(l.count(';') for l in sample_lines) >= sum(l.count(',') for l in sample_lines) else ','
        reader = csv.reader(content.splitlines(), delimiter=delimiter)
        for row in reader:
            if any(cell and str(cell).strip() != '' for cell in row):
                raw_rows.append(row)
    else:
        raise ValueError(f"Desteklenmeyen dosya türü: {ext}")

    if not raw_rows:
        return []

    header_row = [str(c).lower().replace(' ', '').replace('_', '') for c in raw_rows[0]]
    stok_idx, barkod_idx, title_idx, price_idx = 0, 1, 2, 3

    for i, h in enumerate(header_row):
        if any(k in h for k in ['barkod', 'barcode', 'ean', 'gtin']):
            barkod_idx = i
        elif any(k in h for k in ['stokkod', 'urunkod', 'itemcode', 'stockcode']) or h == 'kod':
            stok_idx = i
        elif any(k in h for k in ['malincinsi', 'urunadi', 'stokadi', 'aciklama', 'tanim', 'title', 'product', 'malinadi', 'maladi', 'urun', 'mal']):
            title_idx = i
        elif any(k in h for k in ['satisfiyat', 'fiyat', 'price', 'tutar', 'satfiy', 'sfiyat']):
            price_idx = i

    items = []
    for r in raw_rows[1:]:
        if len(r) <= max(stok_idx, barkod_idx, title_idx, price_idx):
            continue

        raw_code = clean_barcode_text(r[stok_idx]) if stok_idx < len(r) else ""
        raw_barcode = clean_barcode_text(r[barkod_idx]) if barkod_idx < len(r) else ""
        raw_title = normalize_text(r[title_idx]) if title_idx < len(r) else ""
        raw_price = r[price_idx] if price_idx < len(r) else 0

        barcode = raw_barcode or raw_code
        if not barcode or not raw_title:
            continue

        price = parse_price(raw_price)
        parts = raw_title.split(' ', 1)
        brand = parts[0].strip() if len(parts) > 1 else ""

        items.append({
            "barcode": barcode,
            "stock_code": raw_code,
            "title": raw_title,
            "price": price,
            "brand": brand,
            "unit": "ADET"
        })

    return items

def preview_vegawin_comparison(items: list) -> dict:
    """Yüklenen dosyadaki ürünleri ana sistem veritabanı ile karşılaştırır ve detaylı fark listesi üretir."""
    init_db()
    existing_map = {}
    with db_session() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT barcode, title, price, brand, stock_code, label_price FROM urunler;")
        for r in cursor.fetchall():
            existing_map[r["barcode"]] = dict(r)

    comparison_list = []
    price_change_count = 0
    new_product_count = 0
    identical_count = 0
    title_change_count = 0

    for item in items:
        b = item["barcode"]
        inc_p = float(item["price"])
        inc_t = item["title"]
        sc = item.get("stock_code", "")
        brand = item.get("brand", "")

        if b in existing_map:
            main_p = existing_map[b]
            main_price = float(main_p.get("price") or 0)
            main_title = str(main_p.get("title") or "").strip()

            has_price_diff = abs(main_price - inc_p) > 0.001
            has_title_diff = main_title != inc_t

            if has_price_diff:
                diff_amt = round(inc_p - main_price, 2)
                diff_pct = round((diff_amt / main_price * 100) if main_price > 0 else 0, 1)
                status = "price_change"
                price_change_count += 1
            elif has_title_diff:
                diff_amt = 0.0
                diff_pct = 0.0
                status = "title_change"
                title_change_count += 1
            else:
                diff_amt = 0.0
                diff_pct = 0.0
                status = "identical"
                identical_count += 1

            comparison_list.append({
                "barcode": b,
                "stock_code": sc or main_p.get("stock_code", ""),
                "main_title": main_title,
                "incoming_title": inc_t,
                "main_price": main_price,
                "incoming_price": inc_p,
                "diff_amount": diff_amt,
                "diff_percent": diff_pct,
                "status": status,
                "brand": brand or main_p.get("brand", ""),
                "unit": item.get("unit", "ADET")
            })
        else:
            status = "new_product"
            new_product_count += 1
            comparison_list.append({
                "barcode": b,
                "stock_code": sc,
                "main_title": "— (Ana Sistemde Yok)",
                "incoming_title": inc_t,
                "main_price": None,
                "incoming_price": inc_p,
                "diff_amount": 0.0,
                "diff_percent": 0.0,
                "status": status,
                "brand": brand,
                "unit": item.get("unit", "ADET")
            })

    priority = {"price_change": 0, "new_product": 1, "title_change": 2, "identical": 3}
    comparison_list.sort(key=lambda x: (priority.get(x["status"], 99), x["incoming_title"]))

    return {
        "total_incoming": len(items),
        "total_items": len(items),
        "price_changes": price_change_count,
        "price_change_count": price_change_count,
        "new_products": new_product_count,
        "new_product_count": new_product_count,
        "title_changes": title_change_count,
        "title_change_count": title_change_count,
        "identical": identical_count,
        "identical_count": identical_count,
        "items": comparison_list
    }

def sync_vegawin_items(items: list, source_name: str = "VegaWin Aktarımı", device_name: str = "VegaWin PC") -> dict:
    init_db()
    now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    existing_map = {}
    with db_session() as conn:
        cursor = conn.cursor()
        cursor.execute("SELECT barcode, title, price, brand FROM urunler;")
        for r in cursor.fetchall():
            existing_map[r["barcode"]] = dict(r)

    new_count = 0
    updated_count = 0
    price_changes = []
    new_products_list = []

    with db_session() as conn:
        cursor = conn.cursor()
        cursor.execute("""
        INSERT INTO vegawin_sync_history (timestamp, source_file, device_name, total_products, status)
        VALUES (?, ?, ?, ?, 'processing');
        """, (now_str, source_name, device_name, len(items)))
        sync_id = cursor.lastrowid

        for item in items:
            b = item["barcode"]
            p = float(item["price"])
            t = item["title"]
            sc = item.get("stock_code", "")
            brand = item.get("brand", "")

            if b in existing_map:
                old = existing_map[b]
                old_p = float(old.get("price") or 0)
                if abs(old_p - p) > 0.001:
                    diff_amt = round(p - old_p, 2)
                    diff_pct = round((diff_amt / old_p * 100) if old_p > 0 else 0, 1)
                    price_changes.append({
                        "sync_id": sync_id,
                        "barcode": b,
                        "title": t,
                        "old_price": old_p,
                        "new_price": p,
                        "diff_amount": diff_amt,
                        "diff_percent": diff_pct,
                        "changed_at": now_str
                    })
                    cursor.execute("""
                    INSERT INTO vegawin_price_changes (sync_id, barcode, title, old_price, new_price, diff_amount, diff_percent, changed_at, source_device)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?);
                    """, (sync_id, b, t, old_p, p, diff_amt, diff_pct, now_str, device_name))

                    cursor.execute("""
                    UPDATE urunler 
                    SET title = ?, price = ?, stock_code = COALESCE(NULLIF(?, ''), stock_code), 
                        brand = COALESCE(NULLIF(?, ''), brand), updated_at = ?, price_updated_at = ?, 
                        label_price = COALESCE(label_price, ?)
                    WHERE barcode = ?;
                    """, (t, p, sc, brand, now_str, now_str, old_p, b))

                    record_product_history(
                        conn,
                        barcode=b,
                        event_type="price_change",
                        old_title=old.get("title"),
                        new_title=t,
                        old_price=old_p,
                        new_price=p,
                        diff_amount=diff_amt,
                        diff_percent=diff_pct,
                        source="VegaWin Senkronizasyon",
                        device_name=device_name,
                        details=f"Dosya: {source_name}",
                        sync_id=sync_id,
                        timestamp=now_str
                    )
                else:
                    has_title_change = (old.get("title") != t)
                    cursor.execute("""
                    UPDATE urunler 
                    SET title = ?, stock_code = COALESCE(NULLIF(?, ''), stock_code), 
                        brand = COALESCE(NULLIF(?, ''), brand), updated_at = ?
                    WHERE barcode = ?;
                    """, (t, sc, brand, now_str, b))

                    if has_title_change:
                        record_product_history(
                            conn,
                            barcode=b,
                            event_type="title_change",
                            old_title=old.get("title"),
                            new_title=t,
                            source="VegaWin Senkronizasyon",
                            device_name=device_name,
                            details=f"Dosya: {source_name}",
                            sync_id=sync_id,
                            timestamp=now_str
                        )
                updated_count += 1
            else:
                cursor.execute("""
                INSERT INTO urunler (barcode, stock_code, title, price, label_price, brand, is_new, created_at, updated_at, price_updated_at)
                VALUES (?, ?, ?, ?, NULL, ?, 1, ?, ?, ?);
                """, (b, sc, t, p, brand, now_str, now_str, now_str))
                
                cursor.execute("""
                INSERT INTO vegawin_new_products (sync_id, barcode, title, price, created_at, source_device)
                VALUES (?, ?, ?, ?, ?, ?);
                """, (sync_id, b, t, p, now_str, device_name))

                record_product_history(
                    conn,
                    barcode=b,
                    event_type="created",
                    new_title=t,
                    new_price=p,
                    source="VegaWin Senkronizasyon",
                    device_name=device_name,
                    details=f"Yeni Ürün Eklendi (Dosya: {source_name})",
                    sync_id=sync_id,
                    timestamp=now_str
                )
                new_products_list.append({
                    "barcode": b,
                    "title": t,
                    "price": p,
                    "created_at": now_str
                })
                new_count += 1

        cursor.execute("""
        UPDATE vegawin_sync_history 
        SET total_products = ?, new_products = ?, updated_products = ?, price_changes_count = ?, status = 'success'
        WHERE id = ?;
        """, (len(items), new_count, updated_count, len(price_changes), sync_id))

    return {
        "status": "success",
        "sync_id": sync_id,
        "timestamp": now_str,
        "total_received": len(items),
        "new_products": new_count,
        "updated_products": updated_count,
        "price_changes_count": len(price_changes),
        "price_changes": price_changes[:100],
        "new_products_list": new_products_list[:100]
    }

def get_price_changes_list(unprinted_only=False, limit=200):
    with db_session() as conn:
        cursor = conn.cursor()
        if unprinted_only:
            cursor.execute("SELECT * FROM vegawin_price_changes WHERE is_printed = 0 ORDER BY id DESC LIMIT ?;", (limit,))
        else:
            cursor.execute("SELECT * FROM vegawin_price_changes ORDER BY id DESC LIMIT ?;", (limit,))
        return [dict(r) for r in cursor.fetchall()]

def mark_changes_as_printed(change_ids=None):
    with db_session() as conn:
        cursor = conn.cursor()
        if change_ids:
            placeholders = ",".join("?" for _ in change_ids)
            cursor.execute(f"UPDATE vegawin_price_changes SET is_printed = 1 WHERE id IN ({placeholders});", change_ids)
        else:
            cursor.execute("UPDATE vegawin_price_changes SET is_printed = 1 WHERE is_printed = 0;")
